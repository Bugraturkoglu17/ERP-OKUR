"""Fiyat listesi Excel'inden OKUR stok kartlarını okur (veritabanına dokunmaz).

Yalnızca "Stok Kartı ve Fiyatlar" sayfası okunur. Fiyat, görünür "2026 FIYAT" sütunudur
(gizli önerilen/2025/değişim sütunları yok sayılır). İşçilik satırları ("İŞÇİLİKLER", S1..)
hizmettir, stok değildir; atlanır. Fiyatı 0 olan satır "fiyat yok" demektir ve boş saklanır.
"""

from __future__ import annotations

import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

SHEET = "Stok Kartı ve Fiyatlar"
HEADERS = {"code": "MALZEME KODU", "unit": "BİRİM", "category": "ÜRÜN TİPİ", "name": "MALZEME",
           "brand": "MARKA", "spec": "TİP"}
SERVICE_CATEGORY = "İŞÇİLİKLER"


def clean(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return text or None


def code_of(value) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean(value)


def read_catalog(path: Path) -> tuple[list[dict], dict]:
    """(stok kartları, istatistik) döner. Dosya beklenen biçimde değilse ValueError fırlatır."""
    import openpyxl  # yalnızca Excel okunurken gerekir

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if SHEET not in workbook.sheetnames:
            raise ValueError(f"'{SHEET}' sayfası bulunamadı; sayfalar: {workbook.sheetnames}")
        rows = list(workbook[SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()

    header_index = next((i for i, row in enumerate(rows[:10]) if row and clean(row[0]) == HEADERS["code"]), None)
    if header_index is None:
        raise ValueError("Başlık satırı (MALZEME KODU) bulunamadı")
    header = [clean(cell) for cell in rows[header_index]]
    columns = {key: header.index(label) for key, label in HEADERS.items()}
    price_column = next(i for i, label in enumerate(header) if label and label.startswith("2026"))

    items, stats = [], {"services": 0, "zero_price": 0, "no_category": 0}
    for row in rows[header_index + 1:]:
        code = code_of(row[columns["code"]])
        if not code:
            continue
        category = clean(row[columns["category"]])
        if category == SERVICE_CATEGORY:
            stats["services"] += 1
            continue
        price = row[price_column]
        price = Decimal(str(price)).quantize(Decimal("0.0001"), ROUND_HALF_UP) if isinstance(price, (int, float)) else None
        if price is not None and price <= 0:
            price = None
        if price is None:
            stats["zero_price"] += 1
        if category is None:
            stats["no_category"] += 1
        name = clean(row[columns["name"]])
        if not name:
            raise ValueError(f"{code} kodlu satırın adı boş; dosyayı düzeltin")
        items.append({
            "sku": code[:80], "name": name[:255], "category": category[:100] if category else None,
            "unit": (clean(row[columns["unit"]]) or "adet").lower()[:20],
            "brand": clean(row[columns["brand"]]), "spec": clean(row[columns["spec"]]),
            "unit_price": price,
        })
    skus = [item["sku"] for item in items]
    if len(set(skus)) != len(skus):
        raise ValueError("Excel'de tekrar eden malzeme kodu var")
    return items, stats
