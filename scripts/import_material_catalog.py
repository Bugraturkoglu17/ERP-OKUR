"""Import the OKUR material catalog (stok kartı + güncel fiyat) from the price-list Excel.

Reads only the "Stok Kartı ve Fiyatlar" sheet. Price = the visible "2026 FIYAT"
column (the hidden recommended/2025/change columns are ignored). Labour rows
("İŞÇİLİKLER", codes S1..) are services, not stock, and are skipped. A price of 0
means "no price yet" and is stored as empty. Rows are upserted by material code,
so re-running with a newer file updates names/prices and never touches stock
balances or minimum-stock settings.

Needs `pip install openpyxl` (not an application dependency).
    python scripts/import_material_catalog.py "path/to/list.xlsx"           # preview
    python scripts/import_material_catalog.py "path/to/list.xlsx" --apply   # write
"""

from __future__ import annotations

import argparse
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4

import openpyxl
import psycopg2
from psycopg2.extras import execute_values


ROOT = Path(__file__).resolve().parents[1]
ENV_FILE = ROOT / "apps" / "backend" / ".env"
SHEET = "Stok Kartı ve Fiyatlar"
HEADERS = {"code": "MALZEME KODU", "unit": "BİRİM", "category": "ÜRÜN TİPİ", "name": "MALZEME",
           "brand": "MARKA", "spec": "TİP"}
SERVICE_CATEGORY = "İŞÇİLİKLER"


def database_url() -> str:
    values = dict(line.split("=", 1) for line in ENV_FILE.read_text(encoding="utf-8").splitlines()
                  if line and not line.lstrip().startswith("#") and "=" in line)
    url = values["DATABASE_URL"].replace("+asyncpg", "").replace("+psycopg2", "")
    parsed = urlsplit(url)
    if parsed.hostname not in {"localhost", "127.0.0.1"} or parsed.path != "/okur_erp_local":
        raise SystemExit("Refusing to import outside the isolated local OKUR database")
    return url


def clean(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return text or None


def code_of(value) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean(value)


def read_rows(path: Path) -> tuple[list[dict], dict]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise SystemExit(f"Sheet '{SHEET}' not found; sheets: {workbook.sheetnames}")
    sheet = workbook[SHEET]
    header_row = next(i for i, row in enumerate(sheet.iter_rows(max_row=10, values_only=True), 1)
                      if row and clean(row[0]) == HEADERS["code"])
    header = [clean(cell) for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    columns = {key: header.index(label) for key, label in HEADERS.items()}
    price_column = next(i for i, label in enumerate(header) if label and label.startswith("2026"))

    items, stats = [], {"services": 0, "zero_price": 0, "no_category": 0}
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        code = code_of(row[columns["code"]])
        if not code:
            continue
        category = clean(row[columns["category"]])
        if category == SERVICE_CATEGORY:
            stats["services"] += 1
            continue
        price = row[price_column]
        price = Decimal(str(price)).quantize(Decimal("0.0001"), ROUND_HALF_UP) if isinstance(price, (int, float)) else None
        if price is not None and price <= 0:
            price = None
        if price is None:
            stats["zero_price"] += 1
        if category is None:
            stats["no_category"] += 1
        name = clean(row[columns["name"]])
        if not name:
            raise SystemExit(f"Row with code {code} has no name; fix the file first")
        items.append({
            "sku": code[:80], "name": name[:255], "category": category[:100] if category else None,
            "unit": (clean(row[columns["unit"]]) or "adet").lower()[:20],
            "brand": (clean(row[columns["brand"]]) or None), "spec": (clean(row[columns["spec"]]) or None),
            "unit_price": price,
        })
    return items, stats


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", type=Path)
    parser.add_argument("--apply", action="store_true", help="write to the local OKUR database")
    args = parser.parse_args()

    items, stats = read_rows(args.excel)
    skus = [item["sku"] for item in items]
    if len(set(skus)) != len(skus):
        raise SystemExit("Duplicate material codes in the file; fix them first")

    with psycopg2.connect(database_url()) as connection, connection.cursor() as cursor:
        cursor.execute("SELECT id FROM tenants WHERE name = %s", ("Okur Yerel",))
        row = cursor.fetchone()
        if not row:
            raise SystemExit("Tenant 'Okur Yerel' not found")
        tenant_id = row[0]
        cursor.execute("SELECT sku FROM retail_materials WHERE tenant_id = %s", (tenant_id,))
        existing = {r[0] for r in cursor.fetchall()}
        new = [item for item in items if item["sku"] not in existing]
        print(f"File: {len(items)} stock cards | new {len(new)} | update {len(items) - len(new)} | "
              f"skipped services {stats['services']} | no price {stats['zero_price']} | no category {stats['no_category']}")
        if not args.apply:
            print("Preview only. Re-run with --apply to write.")
            return
        execute_values(cursor, """
            INSERT INTO retail_materials (id, tenant_id, sku, name, category, unit, brand, spec, unit_price, min_stock, is_active, created_at)
            VALUES %s
            ON CONFLICT (tenant_id, sku) DO UPDATE SET
                name = EXCLUDED.name, category = EXCLUDED.category, unit = EXCLUDED.unit,
                brand = EXCLUDED.brand, spec = EXCLUDED.spec, unit_price = EXCLUDED.unit_price, is_active = true
        """, [(str(uuid4()), str(tenant_id), i["sku"], i["name"], i["category"], i["unit"], i["brand"], i["spec"],
               i["unit_price"], 0, True) for i in items],
            template="(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, now())", page_size=200)
        cursor.execute("SELECT count(*) FROM retail_materials WHERE tenant_id = %s", (tenant_id,))
        print(f"Applied. Catalog now has {cursor.fetchone()[0]} materials.")


if __name__ == "__main__":
    main()
